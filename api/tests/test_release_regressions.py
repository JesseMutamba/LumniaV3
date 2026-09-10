"""Release regressions for the Lumnia audit repairs.

Copy into api/tests before running with the project's normal environment.
Every HTTP test gets its own temporary SQLite database and test-only keys;
these tests never use conftest's shared database or bootstrap sample data.
"""
from __future__ import annotations

import copy
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook as ExcelWorkbook


@pytest.fixture
def release_api(tmp_path, monkeypatch):
    from app import store
    from app.main import app

    monkeypatch.setattr(store, "DB_PATH", tmp_path / "release.sqlite")
    monkeypatch.setenv("LUMNIA_ADMIN_TOKEN", "release-regression-author-key")
    monkeypatch.setenv("LUMNIA_SESSION_SECRET", "release-regression-session-key")
    monkeypatch.setenv("LUMNIA_RATE_LIMIT", "10000")
    monkeypatch.setenv("LUMNIA_LOGIN_LIMIT", "10000")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    store.init()
    # Do not enter the application lifespan: no bootstrap/seed side effects.
    client = TestClient(app, raise_server_exceptions=False)
    yield client, {"Authorization": "Bearer release-regression-author-key"}
    client.close()


def create_org(api, name="release-a"):
    client, auth = api
    response = client.post(
        "/v1/orgs", headers=auth,
        json={"id": name, "name": name, "sub": {"fr": "Test client"}},
    )
    assert response.status_code == 201
    return name


def report(org, rid="release-report", status="published", amount=123.0):
    return {
        "id": rid, "org": org, "title": {"fr": "Rapport", "en": "Report"},
        "period": {
            "label": {"fr": "T1 2026", "en": "Q1 2026"},
            "start": "2026-01-01", "end": "2026-03-31",
        },
        "status": status,
        "sources": [{"idx": 0, "filename": "source.xlsx"}],
        "blocks": [{
            "type": "kpiGrid", "items": [{
                "label": {"fr": "Revenu", "en": "Revenue"},
                "value": {
                    "n": amount, "unit": "USD",
                    "src": {"file": 0, "sheet": "Sales", "cells": "B2"},
                },
            }],
        }],
    }


def create_reader(api, org, username="release-reader", password="release-test-password"):
    client, auth = api
    created = client.post(
        f"/v1/studio/orgs/{org}/users", headers=auth,
        json={"username": username, "password": password},
    )
    assert created.status_code == 201
    return login(api, username, password)


def login(api, username="release-reader", password="release-test-password"):
    response = api[0].post(
        "/v1/auth/login", json={"username": username, "password": password},
    )
    assert response.status_code == 200
    return {"Authorization": "Bearer " + response.json()["token"]}


def excel_bytes(sheets):
    workbook = ExcelWorkbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def ingest(api, org, documents):
    return api[0].post(
        "/v1/studio/ingest", params={"org": org}, headers=api[1],
        files=[("file", (name, body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
               for name, body in documents],
    )


def test_report_id_cannot_move_between_clients_or_change_an_old_shared_link(release_api):
    client, auth = release_api
    a, b = create_org(release_api), create_org(release_api, "release-b")
    original = client.post(f"/v1/orgs/{a}/reports", headers=auth, json=report(a)).json()
    replacement = client.post(
        f"/v1/orgs/{b}/reports", headers=auth, json=report(b, amount=999.0),
    )
    assert replacement.status_code == 409
    old_link = client.get("/v1/reports/release-report", params={"k": original["share_key"]})
    assert old_link.status_code == 200
    assert old_link.json()["org"] == a
    assert old_link.json()["blocks"][0]["items"][0]["value"]["n"] == 123.0
    assert client.get(f"/v1/orgs/{b}/reports", headers=auth).json() == []


def test_password_reset_invalidates_an_already_issued_session(release_api):
    client, auth = release_api
    org = create_org(release_api)
    old_session = create_reader(release_api, org)
    assert client.get("/v1/auth/me", headers=old_session).status_code == 200
    reset = client.put(
        "/v1/studio/users/release-reader/password", headers=auth,
        json={"password": "release-new-test-password"},
    )
    assert reset.status_code == 200
    assert client.get("/v1/auth/me", headers=old_session).status_code == 401
    current = login(release_api, password="release-new-test-password")
    assert client.get("/v1/auth/me", headers=current).status_code == 200


def test_reenabling_an_account_does_not_revive_its_old_session(release_api):
    client, auth = release_api
    old_session = create_reader(release_api, create_org(release_api))
    path = "/v1/studio/users/release-reader/disable"
    assert client.post(path, headers=auth, params={"disabled": True}).status_code == 200
    assert client.get("/v1/auth/me", headers=old_session).status_code == 401
    assert client.post(path, headers=auth, params={"disabled": False}).status_code == 200
    assert client.get("/v1/auth/me", headers=old_session).status_code == 401
    assert client.get("/v1/auth/me", headers=login(release_api)).status_code == 200


def test_reissuing_a_username_to_another_client_cannot_revive_old_access(release_api):
    client, auth = release_api
    a, b = create_org(release_api), create_org(release_api, "release-b")
    old_session = create_reader(release_api, a)
    assert client.delete("/v1/studio/users/release-reader", headers=auth).status_code == 204
    new_session = create_reader(release_api, b, password="release-reissued-password")
    assert client.get("/v1/auth/me", headers=old_session).status_code == 401
    current = client.get("/v1/auth/me", headers=new_session)
    assert current.status_code == 200 and current.json()["org"]["id"] == b


def test_draft_is_hidden_from_both_keyed_and_authenticated_readers(release_api):
    client, auth = release_api
    org = create_org(release_api)
    reader = create_reader(release_api, org)
    draft = client.post(
        f"/v1/orgs/{org}/reports", headers=auth, json=report(org, status="draft"),
    )
    assert draft.status_code == 201 and draft.json()["status"] == "draft"
    path = "/v1/reports/release-report"
    assert client.get(path, params={"k": draft.json()["share_key"]}).status_code == 404
    assert client.get(path, headers=reader).status_code == 404
    assert client.get("/v1/me/reports", headers=reader).json() == []
    assert client.get("/v1/studio/reports/release-report", headers=auth).status_code == 200
    published = client.patch(
        "/v1/studio/reports/release-report/status", headers=auth,
        params={"new_status": "published"},
    )
    assert published.status_code == 200
    assert client.get(path, headers=reader).status_code == 200


@pytest.mark.parametrize("broken_source", [
    {"file": 0, "sheet": "", "cells": "B2"},
    {"file": 0, "sheet": "   ", "cells": "B2"},
    {"file": 0, "sheet": "Sales", "cells": "not-a-cell"},
    {"file": 0, "sheet": "Sales", "cells": "A0"},
    {"file": 0, "sheet": "Sales", "cells": "XFE1"},
    {"file": 0, "sheet": "Sales", "cells": "A1048577"},
    {"file": 999, "sheet": "Sales", "cells": "B2"},
])
def test_invalid_source_is_refused_without_poisoning_the_report_id(release_api, broken_source):
    client, auth = release_api
    org = create_org(release_api)
    bad = report(org)
    bad["blocks"][0]["items"][0]["value"]["src"] = broken_source
    assert client.post(f"/v1/orgs/{org}/reports", headers=auth, json=bad).status_code == 422
    assert client.get("/v1/studio/reports/release-report", headers=auth).status_code == 404
    assert client.post(f"/v1/orgs/{org}/reports", headers=auth, json=report(org)).status_code == 201


def test_source_ids_cannot_claim_an_array_position_that_does_not_exist(release_api):
    """The renderer resolves src.file through sources[src.file], not a search."""
    client, auth = release_api
    org = create_org(release_api)
    bad = report(org)
    bad["sources"][0]["idx"] = 42
    bad["blocks"][0]["items"][0]["value"]["src"]["file"] = 42
    assert client.post(f"/v1/orgs/{org}/reports", headers=auth, json=bad).status_code == 422


@pytest.mark.parametrize("number", ["NaN", "Infinity", "-Infinity"])
def test_nonfinite_report_values_are_rejected_before_storage(release_api, number):
    client, auth = release_api
    org = create_org(release_api)
    bad = report(org)
    bad["blocks"][0]["items"][0]["value"]["n"] = number
    assert client.post(f"/v1/orgs/{org}/reports", headers=auth, json=bad).status_code == 422
    assert client.get("/v1/studio/reports/release-report", headers=auth).status_code == 404


def test_raw_numeric_table_cells_cannot_bypass_provenance(release_api):
    client, auth = release_api
    org = create_org(release_api)
    bad = report(org)
    bad["sources"] = []
    bad["blocks"] = [{
        "type": "table", "columns": [{"key": "amount", "label": {"fr": "Montant"}}],
        "rows": [{"amount": 123.0}],
    }]
    assert client.post(f"/v1/orgs/{org}/reports", headers=auth, json=bad).status_code == 422


@pytest.fixture
def planned_analysis(release_api):
    client, auth = release_api
    org = create_org(release_api)
    context = {
        "modules": ["budget-vs-actual"],
        "metrics": {
            name: {"budget": {"sheet": "budget", "label": name},
                   "actual": {"sheet": "actual", "label": name}}
            for name in ("OPEX", "CAPEX")
        },
    }
    assert client.put(f"/v1/studio/orgs/{org}/context", headers=auth, json=context).status_code == 200
    sheets = {
        "budget": [["Poste", "Jan", "Fév", "Mar"], ["OPEX", 100, 100, 100], ["CAPEX", 200, 200, 200]],
        "actual": [["Poste", "Jan", "Fév", "Mar"], ["OPEX", 20, 30, 40], ["CAPEX", 50, 60, 70]],
    }
    assert ingest(release_api, org, [("accounts.xlsx", excel_bytes(sheets))]).status_code == 200
    requested = {"org": org, "mode": "analyze", "question": "exécution OPEX"}
    response = client.post("/v1/studio/ask", headers=auth, json=requested)
    assert response.status_code == 200
    plan = response.json()["plan"]
    assert plan["metrics"] == ["OPEX"]
    return org, context, sheets, requested, plan


def execute_plan(api, requested, plan):
    return api[0].post(
        "/v1/studio/ask", headers=api[1],
        json={**requested, "execute": True, "plan": plan},
    )


def test_ask_rejects_a_plan_after_same_filename_gets_new_contents(release_api, planned_analysis):
    from app import store

    org, _, sheets, requested, plan = planned_analysis
    changed = copy.deepcopy(sheets)
    changed["actual"][1][1] = 99
    assert ingest(release_api, org, [("accounts.xlsx", excel_bytes(changed))]).status_code == 200
    runs_before = store.list_runs(org)
    assert execute_plan(release_api, requested, plan).status_code == 409
    assert store.list_runs(org) == runs_before


def test_ask_rejects_a_plan_after_context_changes(release_api, planned_analysis):
    from app import store

    org, context, _, requested, plan = planned_analysis
    revised = {**context, "alert_threshold_pct": 35.0}
    assert release_api[0].put(f"/v1/studio/orgs/{org}/context", headers=release_api[1], json=revised).status_code == 200
    runs_before = store.list_runs(org)
    assert execute_plan(release_api, requested, plan).status_code == 409
    assert store.list_runs(org) == runs_before


def test_ask_executes_only_the_selected_metric(release_api, planned_analysis):
    _, _, _, requested, plan = planned_analysis
    response = execute_plan(release_api, requested, plan)
    assert response.status_code == 200
    payload = response.json()
    metrics = [item.get("metric") for block in payload["blocks"] if block["type"] == "kpiGrid" for item in block["items"]]
    assert metrics == ["OPEX"]
    assert "CAPEX" not in json.dumps(payload["blocks"])
    assert payload["sources"][0]["sha256"] == plan["file_versions"][0]["sha256"]


def test_ask_cannot_execute_without_a_previously_displayed_plan(release_api, planned_analysis):
    _, _, _, requested, _ = planned_analysis
    response = release_api[0].post(
        "/v1/studio/ask", headers=release_api[1], json={**requested, "execute": True},
    )
    assert response.status_code == 409


def test_corrupt_second_workbook_leaves_no_partially_retained_batch(release_api):
    from app import store

    org = create_org(release_api)
    valid = excel_bytes({"Costs": [["Poste", "Budget", "Actual"], ["A", 100, 25], ["B", 200, 30]]})
    response = ingest(release_api, org, [("valid.xlsx", valid), ("broken.xlsx", b"this is not an xlsx file")])
    assert response.status_code == 422
    assert store.latest_files(org) == []
    assert store.list_ingestions(org) == []
    assert store.list_runs(org) == []


def test_failed_analysis_does_not_advance_retained_files_or_recurrence(release_api, monkeypatch):
    from app import store
    from app.routers import ingest as ingest_router

    client, auth = release_api
    org = create_org(release_api)
    assert client.put(f"/v1/studio/orgs/{org}/context", headers=auth, json={"modules": ["execution"]}).status_code == 200
    rows = [["Poste", "Budget", "Actual"], ["A", 100, 25], ["B", 200, 30]]
    assert ingest(release_api, org, [("costs.xlsx", excel_bytes({"Costs": rows}))]).status_code == 200
    files_before = store.latest_files(org)
    baseline_before = store.last_ingestion(org, "costs.xlsx")
    runs_before = store.list_runs(org)

    def unavailable_module(*args, **kwargs):
        raise RuntimeError("Synthetic calculation failure")

    monkeypatch.setattr(ingest_router, "run_modules", unavailable_module)
    changed = copy.deepcopy(rows)
    changed[1][2] = 90
    response = ingest(release_api, org, [("costs.xlsx", excel_bytes({"Costs": changed}))])
    assert response.status_code >= 400
    assert store.latest_files(org) == files_before
    assert store.last_ingestion(org, "costs.xlsx") == baseline_before
    assert store.list_runs(org) == runs_before


def make_snapshot(rows):
    from app.pipeline.ingest import Sheet, Workbook
    from app.pipeline.parse import detect_tables
    from app.pipeline.recur import snapshot
    from app.schema import Source

    workbook = Workbook(
        path=Path("in-memory.xlsx"), sheets={"Data": Sheet("Data", rows)},
        source=Source(idx=0, filename="in-memory.xlsx"),
    )
    return snapshot(workbook, detect_tables(workbook))


def test_recurrence_preserves_duplicate_rows_headers_and_same_sheet_tables():
    from app.pipeline.recur import diff

    rows = [
        ["Label", "Amount", "Amount"], ["Same", 10, 20], ["Same", 30, 40], [None],
        ["Label", "Amount", "Amount"], ["Same", 50, 60], ["Same", 70, 80],
    ]
    previous = make_snapshot(rows)
    values = [v["n"] for table in previous["tables"].values()
              for row in table["rows"].values() for v in row["values"].values()]
    assert len(previous["tables"]) == 2
    assert sorted(values) == [10, 20, 30, 40, 50, 60, 70, 80]
    updated = copy.deepcopy(rows)
    updated[1][1] = 11
    updated[6][1] = 77
    delta = diff(previous, make_snapshot(updated), 20)
    assert {(c["before"], c["after"]) for c in delta["changes"]} == {(10, 11), (70, 77)}


def test_recurrence_does_not_invent_value_changes_when_a_named_row_is_inserted():
    from app.pipeline.recur import diff

    previous = make_snapshot([["Label", "Amount"], ["Alpha", 100], ["Beta", 200]])
    current = make_snapshot([["Label", "Amount"], ["Inserted", 50], ["Alpha", 100], ["Beta", 200]])
    delta = diff(previous, current, 20)
    assert delta["changes"] == []
    assert any("Inserted" in note for note in delta["notes"])


def test_significant_recurrence_alert_is_not_lost_after_many_minor_changes():
    from app.pipeline.recur import MAX_CHANGES, diff

    previous = [["Label", "Amount"]] + [[f"A{i:03d}", 100] for i in range(45)] + [["ZZZ major", 100]]
    current = [["Label", "Amount"]] + [[f"A{i:03d}", 101] for i in range(45)] + [["ZZZ major", 200]]
    delta = diff(make_snapshot(previous), make_snapshot(current), 20)
    assert delta["changes_total"] == 46
    assert delta["alerts_total"] == 1
    assert len(delta["changes"]) == MAX_CHANGES
    assert delta["alerts"][0]["label"] == "ZZZ major"
    assert delta["alerts"][0] in delta["changes"]
    assert delta["notes"]  # Further changes are explicitly disclosed.
